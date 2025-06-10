from docx import Document
import json
import re

def load_template(template_path):
    """Load the template text from file."""
    with open(template_path, 'r') as f:
        return f.read()

def fill_template(template_text, data):
    """Fill in the template with the provided data."""
    # Create a copy of the template text
    filled_text = template_text
    
    # Replace all placeholders with their corresponding values
    for key, value in data.items():
        placeholder = f"{{{key}}}"
        filled_text = filled_text.replace(placeholder, str(value))
    
    return filled_text

def create_word_document(content, output_path):
    """Create a Word document with the filled content."""
    doc = Document()
    
    # Split content into paragraphs and add them to the document
    paragraphs = content.split('\n')
    for para in paragraphs:
        if para.strip():  # Only add non-empty paragraphs
            doc.add_paragraph(para)
    
    # Save the document
    doc.save(output_path)

def main():
    # Load the extracted data (assuming it's saved as JSON)
    try:
        with open('extracted_data.json', 'r') as f:
            data = json.load(f)
    except FileNotFoundError:
        print("Error: extracted_data.json not found. Please run field_extractor.py first.")
        return

    # Load and fill the template
    template_text = load_template('template.txt')
    filled_content = fill_template(template_text, data)

    # Create the Word document
    output_path = 'offer_letter.docx'
    create_word_document(filled_content, output_path)
    print(f"Offer letter has been generated successfully: {output_path}")

if __name__ == "__main__":
    main() 